from datetime import date
from decimal import Decimal
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from unittest.mock import Mock

import requests

from openpyxl import Workbook, load_workbook
from PIL import Image

import receiptcheck as rc


def sample(**changes):
    values = dict(vendor="Sample Mart", date=date(2026, 10, 18), currency="PKR",
                  subtotal="600", tax="60", discount="0", total="660", notes=[],
                  items=[dict(description="Notebook", quantity="2", unit_price="250", line_total="500"),
                         dict(description="Pen", quantity="1", unit_price="100", line_total="100")])
    values.update(changes)
    return rc.Receipt.model_validate(values)


class ReceiptTests(unittest.TestCase):
    def test_extraction_schema_validation_and_connection_errors(self):
        response = Mock()
        response.json.return_value = {"message": {"content": sample().model_dump_json()}}
        with patch.object(rc.requests, "post", return_value=response):
            self.assertEqual(rc.extract_receipt(b"image").total, 660)
            response.json.return_value = {"message": {"content": '{"vendor": "invented"}'}}
            with self.assertRaises(rc.ReceiptError):
                rc.extract_receipt(b"image")
            response.json.return_value = {"done_reason": "length", "message": {"content": "{}"}}
            with self.assertRaisesRegex(rc.ReceiptError, "too long"):
                rc.extract_receipt(b"image")
        with patch.object(rc.requests, "post", side_effect=requests.Timeout):
            with self.assertRaisesRegex(rc.ReceiptError, "timed out"):
                rc.extract_receipt(b"image")

    def test_arithmetic_and_missing_fields(self):
        self.assertEqual(rc.review_reasons(sample()), [])
        self.assertTrue(any("printed total" in reason for reason in rc.review_reasons(sample(total="760"))))
        self.assertTrue(any("Missing" in reason for reason in rc.review_reasons(sample(vendor=None))))
        self.assertFalse(any("Missing" in reason for reason in rc.review_reasons(sample(total="0"))))
        receipt = sample()
        receipt.items[0].unit_price = Decimal("25")
        self.assertTrue(any("quantity" in reason for reason in rc.review_reasons(receipt)))

    def test_image_validation(self):
        with self.assertRaises(rc.ReceiptError):
            rc.prepare_image(b"not an image")
        output = io.BytesIO()
        Image.new("RGB", (20, 20), "white").save(output, format="PNG")
        self.assertTrue(rc.prepare_image(output.getvalue()).startswith(b"\xff\xd8"))

    def test_transparent_images_get_a_white_background(self):
        output = io.BytesIO()
        Image.new("RGBA", (20, 20), (0, 0, 0, 0)).save(output, format="PNG")
        normalized = Image.open(io.BytesIO(rc.prepare_image(output.getvalue())))
        self.assertEqual(normalized.getpixel((10, 10)), (255, 255, 255))

    def test_excel_preservation_duplicate_and_literal_text(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            workbook = folder / "existing.xlsx"
            original = Workbook()
            original.active.title = "My existing sheet"
            original.active["A1"] = "Keep me"
            original.active["B1"] = "=1+2"
            original.save(workbook)
            with patch.object(rc, "DATA", folder / "data"):
                connected = rc.connect_workbook(str(workbook))
                self.assertEqual(rc.saved_connection(), connected)
                receipt = sample(vendor="=HYPERLINK(\"https://example.com\")")
                self.assertTrue(rc.record_receipt(connected, receipt, b"original", "test.png", b"image")[0])
                self.assertFalse(rc.record_receipt(connected, receipt, b"original", "renamed.png", b"image")[0])
                self.assertEqual(len(rc.read_receipts(connected)), 1)
                book = load_workbook(connected)
                self.assertEqual(book["My existing sheet"]["A1"].value, "Keep me")
                self.assertEqual(book["My existing sheet"]["B1"].value, "=1+2")
                self.assertEqual(book["Receipts"]["D2"].data_type, "s")
                self.assertEqual(book["Receipts"]["J2"].value, 660)
                self.assertEqual(book["Line Items"].max_row, 3)
                book.close()
                self.assertTrue(workbook.with_suffix(".xlsx.bak").is_file())

    def test_conflicting_sheet_is_not_overwritten(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "conflict.xlsx"
            original = Workbook()
            original.active.title = "Receipts"
            original.active["A1"] = "Existing unrelated data"
            original.save(path)
            before = path.read_bytes()
            with patch.object(rc, "DATA", Path(temporary) / "data"):
                with self.assertRaises(rc.ReceiptError):
                    rc.connect_workbook(str(path))
            self.assertEqual(path.read_bytes(), before)

    def test_failed_save_preserves_original(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "receipts.xlsx"
            with patch.object(rc, "DATA", Path(temporary) / "data"):
                rc.connect_workbook(str(path))
                before = path.read_bytes()
                with patch.object(rc.os, "replace", side_effect=PermissionError("Excel has file open")):
                    with self.assertRaises(rc.ReceiptError):
                        rc.record_receipt(path, sample(), b"upload", "a.png", b"image")
                self.assertEqual(path.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
